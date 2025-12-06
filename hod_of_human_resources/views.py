from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.db.models import Count, Sum
from .models import (
    HRDashboard, JobRequisition, HRPolicy, HRNotification, Employee,
    WorkforcePlan, Dispute, SuccessionPlan, Meeting
)
from .forms import (
    JobRequisitionForm, HRPolicyForm, NotificationSettingsForm,
    WorkforcePlanForm, DisputeForm, SuccessionPlanForm, MeetingForm
)
from datetime import datetime, timedelta
import json

@login_required
def dashboard(request):
    """Main HR Manager Dashboard"""
    total_headcount = Employee.objects.count()
    open_requisitions = JobRequisition.objects.filter(status='pending').count()
    
    # These metrics might require more complex calculations or data from other models
    high_priority_cases = 0  # Placeholder
    compliance_alerts = 0  # Placeholder
    turnover_rate = 0.0  # Placeholder
    avg_hire_time = 0  # Placeholder
    satisfaction_score = 0.0  # Placeholder
    
    # Example of budget calculation (replace with your actual logic)
    annual_budget = 1125000000  # Placeholder
    utilized_budget = Employee.objects.aggregate(total_salary=Sum('salary'))['total_salary'] or 0
    budget_utilization = (utilized_budget / annual_budget) * 100 if annual_budget > 0 else 0

    context = {
        'total_headcount': total_headcount,
        'open_requisitions': open_requisitions,
        'high_priority_cases': high_priority_cases,
        'compliance_alerts': compliance_alerts,
        'turnover_rate': turnover_rate,
        'avg_hire_time': avg_hire_time,
        'satisfaction_score': satisfaction_score,
        'budget_utilization': budget_utilization,
        'annual_budget': annual_budget,
        'utilized_budget': utilized_budget
    }
    return render(request, 'hod_of_human_resources/dashboard.html', context)

@login_required
def hr_management(request):
    """HR Management Module"""
    active_tab = request.GET.get('tab', 'workforce')

    total_headcount = Employee.objects.count()
    annual_budget = 1125000000  # Placeholder
    utilized_budget = Employee.objects.aggregate(total_salary=Sum('salary'))['total_salary'] or 0

    # Serialize policies to dictionaries to avoid JSON serialization issues
    policies = HRPolicy.objects.all().order_by('-updated_at')
    policies_data = []
    for policy in policies:
        policies_data.append({
            'id': policy.policy_id,
            'title': policy.title,
            'content': policy.content,
            'version': policy.version,
            'status': policy.status,
            'updated_at': policy.updated_at.isoformat() if policy.updated_at else None,
            'created_by': policy.created_by.username if policy.created_by else None
        })

    context = {
        'active_tab': active_tab,
        'requisitions': JobRequisition.objects.all().order_by('-created_at'),
        'policies': policies_data,
        'total_headcount': total_headcount,
        'forecasted_headcount': total_headcount + 5,  # Example forecast
        'annual_budget': annual_budget,
        'utilized_budget': utilized_budget
    }
    return render(request, 'hod_of_human_resources/hr_management.html', context)

@login_required
def analytics(request):
    """HR Analytics & Reporting"""
    department_headcount = Employee.objects.values('department').annotate(count=Count('id'))
    
    total_employees = Employee.objects.count()
    department_percentages = {}
    for dept in department_headcount:
        percentage = int(round((dept['count'] / total_employees) * 100)) if total_employees > 0 else 0
        department_percentages[dept['department']] = {
            'count': dept['count'],
            'percentage': percentage
        }
        
    annual_budget = 1125000000  # Placeholder
    utilized_budget = Employee.objects.aggregate(total_salary=Sum('salary'))['total_salary'] or 0

    context = {
        'turnover_rate': 0.0,  # Placeholder
        'avg_hire_time': 0,  # Placeholder
        'satisfaction_score': 0.0,  # Placeholder
        'compliance_rate': 0,  # Placeholder
        'annual_budget': annual_budget,
        'utilized_budget': utilized_budget,
        'department_data': department_percentages
    }
    return render(request, 'hod_of_human_resources/analytics.html', context)

@login_required
def financials(request):
    """Financial Management"""
    annual_budget = 1125000000  # Placeholder
    utilized_budget = Employee.objects.aggregate(total_salary=Sum('salary'))['total_salary'] or 0
    
    # Example budget breakdown (replace with your actual logic)
    budget_breakdown = {
        'Salaries & Wages': {'budgeted': 750000000, 'actual': utilized_budget, 'variance': 750000000 - utilized_budget},
        'Recruitment Costs': {'budgeted': 150000000, 'actual': 0, 'variance': 150000000},
        'Training & Development': {'budgeted': 112500000, 'actual': 0, 'variance': 112500000},
        'Employee Wellness': {'budgeted': 112500000, 'actual': 0, 'variance': 112500000}
    }
    
    context = {
        'annual_budget': annual_budget,
        'utilized_budget': utilized_budget,
        'budget_breakdown': budget_breakdown
    }
    return render(request, 'hod_of_human_resources/financials.html', context)

@login_required
def generate_financial_report(request):
    """Generate a new financial report (dummy implementation)"""
    if request.method == 'POST':
        # Dummy report generation logic
        from .models import FinancialReport
        report = FinancialReport.objects.create(report_name='Financial Report - ' + datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        messages.success(request, f'New financial report "{report.report_name}" generated successfully.')
        return redirect('hod_of_human_resources:financials')
    else:
        messages.error(request, 'Invalid request method for generating report.')
        return redirect('hod_of_human_resources:financials')

@login_required
def request_budget_adjustment(request):
    """Handle budget adjustment requests"""
    from .forms import BudgetAdjustmentRequestForm
    if request.method == 'POST':
        form = BudgetAdjustmentRequestForm(request.POST)
        if form.is_valid():
            adjustment_request = form.save(commit=False)
            adjustment_request.requested_by = request.user
            adjustment_request.save()
            messages.success(request, 'Budget adjustment request submitted successfully.')
            return redirect('hod_of_human_resources:financials')
    else:
        form = BudgetAdjustmentRequestForm()
    return render(request, 'hod_of_human_resources/request_budget_adjustment.html', {'form': form})

@login_required
def settings(request):
    """System Settings"""
    return render(request, 'hod_of_human_resources/settings.html')

@login_required
def help_support(request):
    """Help & Support"""
    return render(request, 'hod_of_human_resources/help.html')

# API Endpoints

@login_required
def api_requisitions(request):
    """API endpoint for job requisitions"""
    requisitions = JobRequisition.objects.all().values()
    return JsonResponse(list(requisitions), safe=False)

@login_required
def create_requisition(request):
    """Create new job requisition"""
    if request.method == 'POST':
        form = JobRequisitionForm(request.POST)
        if form.is_valid():
            requisition = form.save(commit=False)
            requisition.requested_by = request.user
            requisition.save()
            messages.success(request, 'Job requisition created successfully!')
            return redirect('hod_of_human_resources:hr_management')
    else:
        form = JobRequisitionForm()
    
    return render(request, 'hod_of_human_resources/requisition_form.html', {'form': form})

@login_required
def requisition_update(request, requisition_id):
    requisition = get_object_or_404(JobRequisition, requisition_id=requisition_id)
    if request.method == 'POST':
        form = JobRequisitionForm(request.POST, instance=requisition)
        if form.is_valid():
            form.save()
            messages.success(request, 'Job requisition updated successfully!')
            return redirect('hod_of_human_resources:hr_management')
    else:
        form = JobRequisitionForm(instance=requisition)
    return render(request, 'hod_of_human_resources/requisition_form.html', {'form': form})

@login_required
def requisition_delete(request, requisition_id):
    requisition = get_object_or_404(JobRequisition, requisition_id=requisition_id)
    if request.method == 'POST':
        requisition.delete()
        messages.success(request, 'Job requisition deleted successfully!')
        return redirect('hod_of_human_resources:hr_management')
    return render(request, 'hod_of_human_resources/requisition_confirm_delete.html', {'requisition': requisition})

@login_required
def view_policy(request, policy_id):
    """View HR policy details"""
    policy = get_object_or_404(HRPolicy, policy_id=policy_id)
    return render(request, 'hod_of_human_resources/view_policy.html', {'policy': policy})

@login_required
def manage_policy(request, policy_id=None):
    """Manage HR policies"""
    if request.method == 'POST':
        policy_id_post = request.POST.get('policy_id')
        if policy_id_post:
            policy = get_object_or_404(HRPolicy, policy_id=policy_id_post)
        else:
            policy = None
        form = HRPolicyForm(request.POST, request.FILES, instance=policy)
        if form.is_valid():
            form.save()
            messages.success(request, 'Policy saved successfully!')
            return redirect('hod_of_human_resources:hr_management')
    else:
        if policy_id:
            policy = get_object_or_404(HRPolicy, policy_id=policy_id)
        else:
            policy = None
        form = HRPolicyForm(instance=policy)

    return render(request, 'hod_of_human_resources/manage_policy.html', {'form': form, 'policy': policy})

@login_required
def approve_requisition(request, requisition_id):
    """Approve job requisition"""
    requisition = get_object_or_404(JobRequisition, requisition_id=requisition_id)
    requisition.status = 'approved'
    requisition.save()
    messages.success(request, 'Requisition approved successfully!')
    return redirect('hod_of_human_resources:hr_management')

@login_required
def reject_requisition(request, requisition_id):
    """Reject job requisition"""
    requisition = get_object_or_404(JobRequisition, requisition_id=requisition_id)
    requisition.status = 'rejected'
    requisition.save()
    messages.success(request, 'Requisition rejected!')
    return redirect('hod_of_human_resources:hr_management')

@login_required
def payment_page(request):
    """Display the Payment Page for HR to process employee payments."""
    employees = Employee.objects.all().order_by('last_name', 'first_name')

    # Handle CRUD operations
    if request.method == 'POST':
        action = request.POST.get('action')
        employee_id = request.POST.get('employee_id')

        if action == 'create':
            # Create new employee
            Employee.objects.create(
                first_name=request.POST.get('first_name'),
                last_name=request.POST.get('last_name'),
                email=request.POST.get('email'),
                phone=request.POST.get('phone'),
                department=request.POST.get('department'),
                position=request.POST.get('position'),
                salary=request.POST.get('salary'),
                allowance=request.POST.get('allowance', 0),
                overtime=request.POST.get('overtime', 0),
                hire_date=request.POST.get('hire_date'),
                account_number=request.POST.get('account_number'),
                status='active'
            )
            messages.success(request, 'Employee created successfully!')
        elif action == 'update' and employee_id:
            # Update existing employee
            employee = get_object_or_404(Employee, id=employee_id)
            if 'first_name' in request.POST:
                # Full update
                employee.first_name = request.POST.get('first_name')
                employee.last_name = request.POST.get('last_name')
                employee.email = request.POST.get('email')
                employee.phone = request.POST.get('phone')
                employee.department = request.POST.get('department')
                employee.position = request.POST.get('position')
                employee.salary = request.POST.get('salary')
                employee.allowance = request.POST.get('allowance', 0)
                employee.overtime = request.POST.get('overtime', 0)
                employee.hire_date = request.POST.get('hire_date')
                employee.account_number = request.POST.get('account_number')
                messages.success(request, 'Employee updated successfully!')
            else:
                # Salary only update
                employee.salary = request.POST.get('salary')
                messages.success(request, 'Employee salary updated successfully!')
            employee.save()
        elif action == 'update_allowance' and employee_id:
            # Update allowance only
            employee = get_object_or_404(Employee, id=employee_id)
            employee.allowance = request.POST.get('allowance', 0)
            employee.save()
            messages.success(request, 'Employee allowance updated successfully!')
        elif action == 'update_overtime' and employee_id:
            # Update overtime only
            employee = get_object_or_404(Employee, id=employee_id)
            employee.overtime = request.POST.get('overtime', 0)
            employee.save()
            messages.success(request, 'Employee overtime updated successfully!')
        elif action == 'delete' and employee_id:
            # Delete employee
            employee = get_object_or_404(Employee, id=employee_id)
            employee.delete()
            messages.success(request, 'Employee deleted successfully!')

        return redirect('hod_of_human_resources:payment_page')

    return render(request, 'hod_of_human_resources/payment.html', {'employees': employees})

@login_required
def process_payment(request, employee_id):
    """Process payment for an employee with bank API integration and notifications."""
    employee = get_object_or_404(Employee, id=employee_id)
    if request.method == 'POST':
        payment_type = request.POST.get('payment_type')

        # Determine amount based on payment type
        if payment_type == 'salary':
            amount = employee.salary
        elif payment_type == 'allowance':
            amount = employee.allowance
        elif payment_type == 'overtime':
            amount = employee.overtime
        elif payment_type == 'total':
            amount = employee.total_payable
        else:
            amount = employee.salary  # Default fallback

        # Integrate with bank API
        from .bank_api import simulate_bank_transaction
        bank_response = simulate_bank_transaction(employee.full_name, employee.account_number, amount)

        if bank_response['status'] == 'success':
            # Create payment record
            from .models import Payment
            payment = Payment.objects.create(
                employee=employee,
                amount=amount,
                transaction_id=bank_response['transaction_id'],
                status='success'
            )

            # Send notifications
            from .notifications import send_payment_email, send_payment_sms
            email_subject = f"Payment Confirmation - {payment_type.capitalize()}"
            email_message = f"Dear {employee.full_name},\n\nYour {payment_type} payment of TZS {amount:,.0f}/= has been processed successfully.\nTransaction ID: {payment.transaction_id}\n\nThank you."
            sms_message = f"Payment processed: TZS {amount:,.0f}/= for {payment_type}. TXN: {payment.transaction_id}"

            send_payment_email(employee.email, email_subject, email_message)
            send_payment_sms(employee.phone, sms_message)

            messages.success(request, f"{payment_type.capitalize()} payment of TZS {amount:,.0f}/= processed successfully for {employee.full_name}.")
        else:
            # Create failed payment record
            from .models import Payment
            Payment.objects.create(
                employee=employee,
                amount=amount,
                status='failed'
            )
            messages.error(request, f"Payment failed: {bank_response['message']}")
    else:
        messages.error(request, "Invalid request method.")
    return redirect('hod_of_human_resources:payment_page')

@login_required
def export_payment_report_pdf(request):
    """Export payment reports in PDF format."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from django.http import HttpResponse
    from .models import Payment

    # Get payments data
    payments = Payment.objects.select_related('employee').order_by('-payment_date')

    # Create PDF response
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="payment_report.pdf"'

    # Create PDF document
    doc = SimpleDocTemplate(response, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Title
    title = Paragraph("Payment Report", styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 12))

    # Table data
    data = [['Employee', 'Amount', 'Transaction ID', 'Status', 'Date']]
    for payment in payments:
        data.append([
            payment.employee.full_name,
            f"TZS {payment.amount:,.0f}/=",
            payment.transaction_id or 'N/A',
            payment.status.capitalize(),
            payment.payment_date.strftime('%Y-%m-%d %H:%M')
        ])

    # Create table
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))

    elements.append(table)
    doc.build(elements)
    return response

@login_required
def export_payment_report_excel(request):
    """Export payment reports in Excel format."""
    from openpyxl import Workbook
    from django.http import HttpResponse
    from .models import Payment

    # Get payments data
    payments = Payment.objects.select_related('employee').order_by('-payment_date')

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Payment Report"

    # Headers
    headers = ['Employee', 'Amount', 'Transaction ID', 'Status', 'Date']
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)

    # Data
    for row_num, payment in enumerate(payments, 2):
        ws.cell(row=row_num, column=1, value=payment.employee.full_name)
        ws.cell(row=row_num, column=2, value=float(payment.amount))
        ws.cell(row=row_num, column=3, value=payment.transaction_id or 'N/A')
        ws.cell(row=row_num, column=4, value=payment.status.capitalize())
        ws.cell(row=row_num, column=5, value=payment.payment_date.strftime('%Y-%m-%d %H:%M'))

    # Create response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="payment_report.xlsx"'
    wb.save(response)
    return response


# Workforce Planning Views
@login_required
def workforce_plan_list(request):
    plans = WorkforcePlan.objects.all()
    return render(request, 'hod_of_human_resources/workforce_plan_list.html', {'plans': plans})

@login_required
def workforce_plan_create(request):
    if request.method == 'POST':
        form = WorkforcePlanForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Workforce plan created successfully.')
            return redirect('hod_of_human_resources:workforce_plan_list')
    else:
        form = WorkforcePlanForm()
    return render(request, 'hod_of_human_resources/workforce_plan_form.html', {'form': form})

@login_required
def workforce_plan_update(request, pk):
    plan = get_object_or_404(WorkforcePlan, pk=pk)
    if request.method == 'POST':
        form = WorkforcePlanForm(request.POST, instance=plan)
        if form.is_valid():
            form.save()
            messages.success(request, 'Workforce plan updated successfully.')
            return redirect('hod_of_human_resources:workforce_plan_list')
    else:
        form = WorkforcePlanForm(instance=plan)
    return render(request, 'hod_of_human_resources/workforce_plan_form.html', {'form': form})

@login_required
def workforce_plan_delete(request, pk):
    plan = get_object_or_404(WorkforcePlan, pk=pk)
    if request.method == 'POST':
        plan.delete()
        messages.success(request, 'Workforce plan deleted successfully.')
        return redirect('hod_of_human_resources:workforce_plan_list')
    return render(request, 'hod_of_human_resources/workforce_plan_confirm_delete.html', {'plan': plan})


# Dispute Resolution Views
@login_required
def dispute_list(request):
    disputes = Dispute.objects.all()
    return render(request, 'hod_of_human_resources/dispute_list.html', {'disputes': disputes})

@login_required
def dispute_create(request):
    if request.method == 'POST':
        form = DisputeForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Dispute created successfully.')
            return redirect('hod_of_human_resources:dispute_list')
    else:
        form = DisputeForm()
    return render(request, 'hod_of_human_resources/dispute_form.html', {'form': form})

@login_required
def dispute_update(request, pk):
    dispute = get_object_or_404(Dispute, pk=pk)
    if request.method == 'POST':
        form = DisputeForm(request.POST, instance=dispute)
        if form.is_valid():
            form.save()
            messages.success(request, 'Dispute updated successfully.')
            return redirect('hod_of_human_resources:dispute_list')
    else:
        form = DisputeForm(instance=dispute)
    return render(request, 'hod_of_human_resources/dispute_form.html', {'form': form})

@login_required
def dispute_delete(request, pk):
    dispute = get_object_or_404(Dispute, pk=pk)
    if request.method == 'POST':
        dispute.delete()
        messages.success(request, 'Dispute deleted successfully.')
        return redirect('hod_of_human_resources:dispute_list')
    return render(request, 'hod_of_human_resources/dispute_confirm_delete.html', {'dispute': dispute})


# Succession Planning Views
@login_required
def succession_plan_list(request):
    plans = SuccessionPlan.objects.all()
    return render(request, 'hod_of_human_resources/succession_plan_list.html', {'plans': plans})

@login_required
def succession_plan_create(request):
    if request.method == 'POST':
        form = SuccessionPlanForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Succession plan created successfully.')
            return redirect('hod_of_human_resources:succession_plan_list')
    else:
        form = SuccessionPlanForm()
    return render(request, 'hod_of_human_resources/succession_plan_form.html', {'form': form})

@login_required
def succession_plan_update(request, pk):
    plan = get_object_or_404(SuccessionPlan, pk=pk)
    if request.method == 'POST':
        form = SuccessionPlanForm(request.POST, instance=plan)
        if form.is_valid():
            form.save()
            messages.success(request, 'Succession plan updated successfully.')
            return redirect('hod_of_human_resources:succession_plan_list')
    else:
        form = SuccessionPlanForm(instance=plan)
    return render(request, 'hod_of_human_resources/succession_plan_form.html', {'form': form})

@login_required
def succession_plan_delete(request, pk):
    plan = get_object_or_404(SuccessionPlan, pk=pk)
    if request.method == 'POST':
        plan.delete()
        messages.success(request, 'Succession plan deleted successfully.')
        return redirect('hod_of_human_resources:succession_plan_list')
    return render(request, 'hod_of_human_resources/succession_plan_confirm_delete.html', {'plan': plan})


# Meeting Management Views
@login_required
def meeting_list(request):
    meetings = Meeting.objects.all().order_by('-meeting_time')
    return render(request, 'hod_of_human_resources/meeting_list.html', {'meetings': meetings})

@login_required
def meeting_detail(request, pk):
    meeting = get_object_or_404(Meeting, pk=pk)
    return render(request, 'hod_of_human_resources/meeting_detail.html', {'meeting': meeting})

@login_required
def meeting_create(request):
    if request.method == 'POST':
        form = MeetingForm(request.POST)
        if form.is_valid():
            meeting = form.save(commit=False)
            meeting.created_by = request.user
            meeting.save()
            form.save_m2m()  # Save the many-to-many relationships
            
            # Send notifications to participants
            for participant in meeting.participants.all():
                HRNotification.objects.create(
                    title=f"New Meeting Invitation: {meeting.title}",
                    message=f"You have been invited to a meeting on {meeting.meeting_time.strftime('%Y-%m-%d at %H:%M')}. Venue: {meeting.venue}",
                    notification_type='meeting',
                    recipient=participant
                )
            
            messages.success(request, 'Meeting scheduled successfully and participants notified.')
            return redirect('hod_of_human_resources:meeting_list')
    else:
        form = MeetingForm()
    return render(request, 'hod_of_human_resources/meeting_form.html', {'form': form})

@login_required
def meeting_update(request, pk):
    meeting = get_object_or_404(Meeting, pk=pk)
    if request.method == 'POST':
        form = MeetingForm(request.POST, instance=meeting)
        if form.is_valid():
            form.save()
            messages.success(request, 'Meeting updated successfully.')
            return redirect('hod_of_human_resources:meeting_list')
    else:
        form = MeetingForm(instance=meeting)
    return render(request, 'hod_of_human_resources/meeting_form.html', {'form': form})

@login_required
def meeting_delete(request, pk):
    meeting = get_object_or_404(Meeting, pk=pk)
    if request.method == 'POST':
        meeting.delete()
        messages.success(request, 'Meeting deleted successfully.')
        return redirect('hod_of_human_resources:meeting_list')
    return render(request, 'hod_of_human_resources/meeting_confirm_delete.html', {'meeting': meeting})
